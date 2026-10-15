import io

from docx import Document

from tests.conftest import NOTES_MD, register


def upload_md(client, auth, **extra):
    files = {"file": ("конспект.md", NOTES_MD.encode(), "text/markdown")}
    return client.post("/materials", files=files, data=extra, headers=auth)


# ---------- auth ----------
def test_register_login_me_and_duplicates(client):
    auth = register(client)
    me = client.get("/auth/me", headers=auth).json()
    assert me["username"] == "vlad" and "password" not in me and "password_hash" not in me
    dup = client.post("/auth/register",
                      json={"email": "VLAD@example.com", "username": "other", "password": "secret123"})
    assert dup.status_code == 409
    assert client.post("/auth/login", json={"login": "vlad", "password": "wrong"}).status_code == 401
    # вход по email
    assert client.post("/auth/login",
                       json={"login": "vlad@example.com", "password": "secret123"}).status_code == 200
    assert client.post("/auth/token", data={"username": "vlad", "password": "secret123"}).status_code == 200


def test_protected_routes_need_token(client):
    assert client.get("/decks").status_code == 401
    assert client.get("/decks", headers={"Authorization": "Bearer nonsense"}).status_code == 401


def test_demo_deck_is_public(client):
    r = client.get("/demo/deck")
    assert r.status_code == 200 and len(r.json()["cards"]) >= 3


# ---------- disciplines / decks / cards ----------
def test_discipline_and_deck_crud(client, auth):
    d = client.post("/disciplines", json={"name": "Математика"}, headers=auth).json()
    deck = client.post("/decks", json={"title": "Пределы", "discipline_id": d["id"]}, headers=auth).json()
    assert deck["card_count"] == 0

    card = client.post(f"/decks/{deck['id']}/cards",
                       json={"question": "2+2?", "answer": "4"}, headers=auth).json()
    patched = client.patch(f"/cards/{card['id']}", json={"answer": "четыре"}, headers=auth).json()
    assert patched["answer"] == "четыре"

    listed = client.get("/decks", params={"discipline_id": d["id"]}, headers=auth).json()
    assert len(listed) == 1 and listed[0]["card_count"] == 1 and listed[0]["due_count"] == 1
    assert client.get("/decks", params={"discipline_id": 9999}, headers=auth).json() == []

    # снять дисциплину явным null
    r = client.patch(f"/decks/{deck['id']}", json={"discipline_id": None}, headers=auth).json()
    assert r["discipline_id"] is None

    # удаление дисциплины не удаляет колоду
    client.patch(f"/decks/{deck['id']}", json={"discipline_id": d["id"]}, headers=auth)
    assert client.delete(f"/disciplines/{d['id']}", headers=auth).status_code == 204
    assert client.get(f"/decks/{deck['id']}", headers=auth).json()["discipline_id"] is None

    # каскад: удаление колоды удаляет карточки
    assert client.delete(f"/decks/{deck['id']}", headers=auth).status_code == 204
    assert client.get(f"/decks/{deck['id']}", headers=auth).status_code == 404
    assert client.delete(f"/cards/{card['id']}", headers=auth).status_code == 404


def test_users_cannot_see_each_others_data(client):
    a = register(client, "alice", "a@example.com")
    b = register(client, "bob", "b@example.com")
    deck = client.post("/decks", json={"title": "Секрет"}, headers=a).json()
    assert client.get(f"/decks/{deck['id']}", headers=b).status_code == 404
    assert client.delete(f"/decks/{deck['id']}", headers=b).status_code == 404
    assert client.get("/decks", headers=b).json() == []
    assert client.post("/decks", json={"title": "x", "discipline_id": 1}, headers=b).status_code == 404


# ---------- materials ----------
def test_upload_markdown_creates_deck(client, auth):
    d = client.post("/disciplines", json={"name": "Педагогика"}, headers=auth).json()
    r = upload_md(client, auth, discipline_id=d["id"], title="Интервальное обучение")
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["material"]["status"] == "processed" and body["material"]["processor_type"] == "heuristic"
    assert body["deck"]["discipline_id"] == d["id"] and body["deck"]["card_count"] >= 3
    assert len(client.get("/materials", headers=auth).json()) == 1


def test_upload_docx(client, auth):
    doc = Document()
    doc.add_heading("Термины", level=1)
    doc.add_paragraph("Колода — набор карточек, объединённых одной темой или конспектом.")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/materials", headers=auth,
                    files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    assert r.json()["material"]["file_format"] == "docx"


def test_upload_rejects_bad_files(client, auth):
    bad_ext = client.post("/materials", headers=auth, files={"file": ("a.exe", b"zzz", "x/y")})
    assert bad_ext.status_code == 415
    no_text = client.post("/materials", headers=auth, files={"file": ("a.md", b"hi", "text/plain")})
    assert no_text.status_code == 422
    assert client.get("/materials", headers=auth).json()[0]["status"] == "failed"
    garbage_pdf = client.post("/materials", headers=auth, files={"file": ("a.pdf", b"not pdf", "x/y")})
    assert garbage_pdf.status_code == 422


# ---------- study ----------
def test_full_study_flow_and_stats(client, auth):
    # статистика до первой сессии не показывается
    ov = client.get("/stats/overview", headers=auth).json()
    assert ov["has_data"] is False and ov["accuracy_percent"] is None
    assert client.get("/stats/weekly", headers=auth).json()["has_data"] is False

    deck = upload_md(client, auth).json()["deck"]
    session = client.post(f"/decks/{deck['id']}/sessions", headers=auth).json()

    answered, first_wrong = 0, None
    last_id = None
    while True:
        nxt = client.get(f"/sessions/{session['id']}/next", headers=auth).json()
        card = nxt["card"]
        if card is None:
            break
        assert card["id"] != last_id or nxt["remaining_due"] == 1  # подряд одну и ту же не показываем
        correct = not (first_wrong is None)  # первый ответ — неверный
        if first_wrong is None:
            first_wrong = card["id"]
        ans = client.post(f"/sessions/{session['id']}/answers", headers=auth, json={
            "card_id": card["id"], "is_correct": correct, "response_time_ms": 2000})
        assert ans.status_code == 201, ans.text
        answered += 1
        last_id = card["id"]
        assert answered < 50, "сессия не завершается"

    # неверно отвеченная карточка вернулась и была отвечена верно
    results_total = deck["card_count"] + 1
    assert answered == results_total

    fin = client.post(f"/sessions/{session['id']}/finish", headers=auth).json()
    assert fin["finished_at"] is not None and fin["total_answers"] == answered
    assert client.post(f"/sessions/{session['id']}/finish", headers=auth).status_code == 200
    assert client.get(f"/sessions/{session['id']}/next", headers=auth).status_code == 409
    assert client.post(f"/sessions/{session['id']}/answers", headers=auth,
                       json={"card_id": first_wrong, "is_correct": True}).status_code == 409

    ov = client.get("/stats/overview", headers=auth).json()
    assert ov["has_data"] and ov["total_answers"] == answered and ov["total_sessions"] == 1
    assert ov["accuracy_percent"] == round(100 * (answered - 1) / answered, 1)
    assert ov["due_cards"] == 0

    week = client.get("/stats/weekly", headers=auth).json()
    assert week["has_data"] and len(week["daily_activity"]) == 7
    assert week["learned_cards"] == deck["card_count"]
    assert week["total_study_seconds"] == 2 * answered
    assert sum(d["cards_answered"] for d in week["daily_activity"]) == answered


def test_answer_validation(client, auth):
    deck = client.post("/decks", json={"title": "A"}, headers=auth).json()
    other = client.post("/decks", json={"title": "B"}, headers=auth).json()
    card_b = client.post(f"/decks/{other['id']}/cards", json={"question": "q", "answer": "a"},
                         headers=auth).json()
    s = client.post(f"/decks/{deck['id']}/sessions", headers=auth).json()
    r = client.post(f"/sessions/{s['id']}/answers", headers=auth,
                    json={"card_id": card_b["id"], "is_correct": True})
    assert r.status_code == 400
    r = client.post(f"/sessions/{s['id']}/answers", headers=auth,
                    json={"card_id": card_b["id"], "is_correct": True, "quality": 9})
    assert r.status_code == 422
    # пустая колода: повторять нечего
    assert client.get(f"/sessions/{s['id']}/next", headers=auth).json() == {"card": None, "remaining_due": 0}


# ---------- notifications ----------
def test_notifications_and_reminder(client, auth):
    cur = client.get("/notifications", headers=auth).json()
    assert cur["enabled"] is False and cur["days"] == [0, 1, 2, 3, 4, 5, 6]

    upd = client.put("/notifications", headers=auth,
                     json={"enabled": True, "notify_time": "08:30", "days": [4, 2, 2]}).json()
    assert upd == {"enabled": True, "notify_time": "08:30", "days": [2, 4]}
    assert client.put("/notifications", headers=auth,
                      json={"enabled": True, "notify_time": "25:99"}).status_code == 422
    assert client.put("/notifications", headers=auth,
                      json={"enabled": True, "days": [7]}).status_code == 422

    upload_md(client, auth)  # есть карточки к повторению
    client.put("/notifications", headers=auth,
               json={"enabled": True, "notify_time": "08:30", "days": [0, 1, 2, 3, 4, 5, 6]})
    r = client.get("/notifications/reminder", headers=auth).json()
    assert r["should_remind"] is True and r["due_cards"] >= 3
    client.put("/notifications", headers=auth, json={"enabled": False})
    assert client.get("/notifications/reminder", headers=auth).json()["should_remind"] is False


# ---------- account ----------
def test_delete_account_removes_everything(client):
    auth = register(client)
    upload_md(client, auth)
    assert client.delete("/auth/me", headers=auth).status_code == 204
    assert client.get("/auth/me", headers=auth).status_code == 401
    # логин с теми же данными больше невозможен
    assert client.post("/auth/login", json={"login": "vlad", "password": "secret123"}).status_code == 401
