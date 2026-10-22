"""DocumentProcessor из UML: извлечение текста из PDF, DOCX, DOC и MD."""
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

SUPPORTED_FORMATS = {"pdf", "docx", "doc", "md"}


class DocumentProcessingError(Exception):
    """Файл нельзя обработать (формат, повреждение, нет текста)."""


def detect_format(file_name: str) -> str:
    ext = Path(file_name).suffix.lower().lstrip(".")
    if ext == "markdown":
        ext = "md"
    if ext not in SUPPORTED_FORMATS:
        raise DocumentProcessingError(
            f"Формат «.{ext or '?'}» не поддерживается. Допустимы: PDF, DOCX, DOC, MD.")
    return ext


def extract_text(file_name: str, data: bytes) -> tuple[str, str]:
    """Возвращает (текст, формат). Заголовки DOCX/MD помечаются как «# …»."""
    fmt = detect_format(file_name)
    if fmt == "pdf":
        text = _pdf_text(data)
    elif fmt == "docx":
        text = _docx_text(data)
    elif fmt == "doc":
        text = _doc_text(data)
    else:
        text = _decode(data)
    text = _normalize(text)
    if len(text) < 20:
        raise DocumentProcessingError("В файле не найден текст для создания карточек.")
    return text, fmt


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except (PyPdfError, ValueError, OSError) as exc:
        raise DocumentProcessingError(f"Не удалось прочитать PDF: {exc}") from exc
    text = "\n\n".join(pages)
    if not text.strip():
        raise DocumentProcessingError(
            "В PDF нет текстового слоя (похоже на скан). Загрузите PDF с выделяемым текстом.")
    return text


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx бросает разные исключения на битых файлах
        raise DocumentProcessingError(f"Не удалось прочитать DOCX: {exc}") from exc
    lines: list[str] = []
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith(("heading", "заголовок", "title")):
            lines.append(f"# {txt}")
        else:
            lines.append(txt)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if len(cells) == 2:
                lines.append(f"{cells[0]} — {cells[1]}")
            elif cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)


def _doc_text(data: bytes) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        antiword = shutil.which("antiword")
        if antiword:
            proc = subprocess.run([antiword, str(src)], capture_output=True, timeout=30)
            if proc.returncode == 0 and proc.stdout.strip():
                return proc.stdout.decode("utf-8", errors="replace")
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
                capture_output=True, timeout=90)
            converted = Path(tmp) / "input.docx"
            if converted.exists():
                return _docx_text(converted.read_bytes())
    raise DocumentProcessingError(
        "Для старого формата .doc на сервере нужен antiword или LibreOffice. "
        "Сохраните файл как .docx и загрузите снова.")
